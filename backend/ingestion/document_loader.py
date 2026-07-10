from langchain_community.document_loaders import PyMuPDFLoader, Docx2txtLoader, CSVLoader
from langchain_core.documents import Document
from openpyxl import load_workbook
from pptx import Presentation
import os

LOADERS = {
    ".pdf": PyMuPDFLoader,
    ".docx": Docx2txtLoader,
    ".csv": CSVLoader,
}


def _load_xlsx(file_path: str, doc_id: str) -> list[Document]:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    pages = []

    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                rows.append(" | ".join(values))

        if rows:
            pages.append(
                Document(
                    page_content="\n".join(rows),
                    metadata={
                        "doc_id": doc_id,
                        "doc_name": os.path.basename(file_path),
                        "file_type": ".xlsx",
                        "source_type": "document",
                        "page_num": None,
                        "sheet_name": sheet.title,
                        "source": os.path.basename(file_path),
                    },
                )
            )

    workbook.close()
    return pages


def _load_pptx(file_path: str, doc_id: str) -> list[Document]:
    presentation = Presentation(file_path)
    pages = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                text = shape.text.strip()
                if text:
                    parts.append(text)

        if parts:
            pages.append(
                Document(
                    page_content="\n".join(parts),
                    metadata={
                        "doc_id": doc_id,
                        "doc_name": os.path.basename(file_path),
                        "file_type": ".pptx",
                        "source_type": "document",
                        "page_num": slide_index,
                        "slide_index": slide_index,
                        "source": os.path.basename(file_path),
                    },
                )
            )

    return pages


def load_document(file_path: str, doc_id: str) -> list:
    ext = os.path.splitext(file_path)[-1].lower()

    if ext == ".xlsx":
        pages = _load_xlsx(file_path, doc_id)
    elif ext == ".pptx":
        pages = _load_pptx(file_path, doc_id)
    elif ext in LOADERS:
        loader = LOADERS[ext](file_path)
        pages = loader.load()
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    for i, page in enumerate(pages):
        page.metadata["doc_id"] = doc_id
        page.metadata["doc_name"] = os.path.basename(file_path)  # add explicitly
        page.metadata["file_type"] = ext
        page.metadata["source_type"] = "document"

        if ext == ".pdf":
            page.metadata["page_num"] = page.metadata.get("page", i) + 1  # fix 0-based
        elif ext in [".csv", ".xlsx"]:
            page.metadata["page_num"] = None
            page.metadata["row_index"] = i
        elif ext == ".docx":
            page.metadata["page_num"] = None
            page.metadata["section_index"] = i
        elif ext == ".pptx":
            page.metadata["page_num"] = None
            page.metadata["slide_index"] = i

        page.metadata["source"] = os.path.basename(file_path)

    if not pages:
        raise ValueError(f"No readable text found in {os.path.basename(file_path)}")

    return pages
